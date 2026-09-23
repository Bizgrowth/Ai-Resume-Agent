import os
import io
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from resume_agent import ResumeAgentPipeline, OptimizedResume

st.set_page_config(
    page_title="AI Resume Optimizer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="collapsed"
)

st.markdown("""
<style>
    @media (max-width: 767px) {
        .block-container {
            padding: 1rem 0.6rem !important;
        }
    }
    .stButton>button {
        width: 100%;
        height: 3.2rem;
        font-size: 1.05rem;
        font-weight: 600;
        border-radius: 8px;
    }
</style>
""", unsafe_allow_html=True)

def generate_docx_bytes(resume: OptimizedResume) -> io.BytesIO:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    def add_heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(resume.full_name)
    name_run.bold = True
    name_run.font.size = Pt(18)
    name_run.font.name = "Calibri"

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(12)
    c_run = contact_p.add_run(resume.contact_info)
    c_run.font.size = Pt(9.5)
    c_run.font.name = "Calibri"

    add_heading("Professional Summary")
    sum_p = doc.add_paragraph()
    sum_p.paragraph_format.space_after = Pt(8)
    sum_run = sum_p.add_run(resume.professional_summary)
    sum_run.font.size = Pt(10)
    sum_run.font.name = "Calibri"

    add_heading("Core Competencies & Keywords")
    comp_p = doc.add_paragraph()
    comp_p.paragraph_format.space_after = Pt(8)
    comp_run = comp_p.add_run(" • ".join(resume.core_competencies))
    comp_run.font.size = Pt(9.5)
    comp_run.font.name = "Calibri"

    if resume.technical_stack:
        add_heading("Technical & Operational Stack")
        stack_p = doc.add_paragraph()
        stack_p.paragraph_format.space_after = Pt(8)
        stack_run = stack_p.add_run(" • ".join(resume.technical_stack))
        stack_run.font.size = Pt(9.5)
        stack_run.font.name = "Calibri"

    add_heading("Professional Experience")
    for role in resume.experience:
        role_header = doc.add_paragraph()
        role_header.paragraph_format.space_before = Pt(6)
        role_header.paragraph_format.space_after = Pt(2)
        role_header.paragraph_format.keep_with_next = True
        
        t_run = role_header.add_run(f"{role.role_title} | {role.company}")
        t_run.bold = True
        t_run.font.size = Pt(10.5)
        t_run.font.name = "Calibri"
        
        d_run = role_header.add_run(f"    ({role.dates})")
        d_run.italic = True
        d_run.font.size = Pt(9.5)
        d_run.font.name = "Calibri"

        for bullet in role.bullets:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(2)
            bp.paragraph_format.space_before = Pt(0)
            brun = bp.add_run(bullet)
            brun.font.size = Pt(9.5)
            brun.font.name = "Calibri"

    if resume.education:
        add_heading("Education & Credentials")
        for edu in resume.education:
            ed_p = doc.add_paragraph(style='List Bullet')
            ed_p.paragraph_format.space_after = Pt(2)
            ed_run = ed_p.add_run(edu)
            ed_run.font.size = Pt(9.5)
            ed_run.font.name = "Calibri"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

st.title("📄 AI Resume Optimizer")
st.caption("Cross-Platform ATS Alignment Engine")

if "optimized_resume" not in st.session_state:
    st.session_state.optimized_resume = None
if "docx_buffer" not in st.session_state:
    st.session_state.docx_buffer = None

api_key = os.getenv("GEMINI_API_KEY")
if not api_key:
    with st.expander("⚙️ System Credentials (API Key Required)"):
        api_key = st.text_input("Gemini API Key", type="password")

col1, col2 = st.columns([1, 1], gap="large")

with col1:
    st.subheader("📥 Input Data")
    raw_jd = st.text_area("1. Target Job Description", placeholder="Paste job posting here...", height=180)
    raw_resume = st.text_area("2. Candidate Work History", placeholder="Paste resume or messy achievements here...", height=200)
    generate_btn = st.button("🚀 Optimize Resume for ATS", type="primary")

with col2:
    st.subheader("📤 Output & ATS Evaluation")
    if generate_btn:
        if not api_key:
            st.error("Please provide your Gemini API Key.")
        elif not raw_jd.strip() or not raw_resume.strip():
            st.warning("Please paste both the Job Description and Candidate Data.")
        else:
            with st.status("Optimizing resume...", expanded=True) as status:
                try:
                    agent = ResumeAgentPipeline(api_key=api_key)
                    resume_data = agent.run(raw_resume=raw_resume, raw_jd=raw_jd)
                    docx_file = generate_docx_bytes(resume_data)
                    
                    st.session_state.optimized_resume = resume_data
                    st.session_state.docx_buffer = docx_file
                    status.update(label="Complete!", state="complete", expanded=False)
                    st.success("Resume generated and ATS aligned!")
                except Exception as e:
                    status.update(label="Failed", state="error")
                    st.error(f"Error: {str(e)}")

    if st.session_state.optimized_resume:
        res = st.session_state.optimized_resume
        file_name = f"{res.full_name.replace(' ', '_')}_ATS_Resume.docx"
        
        st.download_button(
            label="📥 Download ATS Word Document (.docx)",
            data=st.session_state.docx_buffer,
            file_name=file_name,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

        with st.expander("👀 View ATS Preview", expanded=True):
            st.markdown(f"### {res.full_name}")
            st.caption(res.contact_info)
            st.markdown("**Summary:**")
            st.write(res.professional_summary)
            st.markdown("**Keywords Matched:**")
            st.write(" • ".join(res.core_competencies))
            st.markdown("**Experience:**")
            for r in res.experience:
                st.markdown(f"**{r.role_title}** | *{r.company}* ({r.dates})")
                for b in r.bullets:
                    st.write(f"- {b}")
                  
